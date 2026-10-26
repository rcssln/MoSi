"""report.py - presentation outputs: charts (PNG) and one Excel workbook.

Run after run_experiments.py, analyze.py and sensitivity.py.
Needs openpyxl (only for the Excel file).
"""
import os
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import config
from analyze import MEASURES, summarize

OUT = config.OUTPUT_DIR
COLORS = {"baseline": "#4C72B0", "timed": "#DD8452"}
LABELS = {"baseline": "Baseline (leave when full)",
          "timed": f"Timed (leave when full or after T = {config.TIME_LIMITS[0]} min)"}
PANELS = ["avg_wait", "avg_depart_wait", "max_depart_wait", "avg_idle", "avg_seats",
          "served_per_hour"]


def plot_key_measures(summ):
    """Grouped bars: baseline vs timed, peak vs off-peak, 95% CI as error bars."""
    fig, axes = plt.subplots(2, 3, figsize=(14, 8))
    for ax, m in zip(axes.flat, PANELS):
        for i, rule in enumerate(("baseline", "timed")):
            g = summ[summ.rule == rule].sort_values("period", ascending=False)  # peak first
            x = [j + (i - 0.5) * 0.35 for j in range(len(g))]
            ax.bar(x, g[m + "_mean"], 0.35, yerr=g[m + "_hw"], capsize=4,
                   color=COLORS[rule], label=LABELS[rule])
        ax.set_xticks(range(2))
        ax.set_xticklabels(["Peak", "Off-peak"])
        ax.set_ylabel(MEASURES[m])
        ax.set_title(MEASURES[m])
    handles, labels = axes[0, 0].get_legend_handles_labels()
    fig.legend(handles, labels, loc="lower center", ncol=2, fontsize=10)
    fig.suptitle(f"Baseline vs timed dispatch, N = {config.FLEET_SIZES[0]} jeepneys "
                 "(bars = mean, whiskers = 95% CI)")
    fig.tight_layout(rect=(0, 0.05, 1, 1))
    fig.savefig(os.path.join(OUT, "key_measures.png"), dpi=150)
    plt.close(fig)


def capacity_per_min(n):
    """Rough max passengers/min the fleet can carry: N x C per (round trip + loading)."""
    load = config.CAPACITY * config.BOARD_MEAN_SEC / 60
    return n * config.CAPACITY / (config.TRIP_MEAN_MIN["peak"] + load)


def plot_sweep(df, xcol, xlabel, title, fname, vline=None, vlabel=None):
    panels = [("avg_wait", "Average wait in line (min)"),
              ("avg_depart_wait", "Average time from arrival to departure (min)"),
              ("avg_idle", "Jeepney time at terminal per trip (min)"),
              ("avg_seats", "Seats filled per departure (people)"),
              ("served_per_hour", "Passengers served (per hour)")]
    fig, axes = plt.subplots(2, 3, figsize=(15, 8))
    axes.flat[-1].axis("off")
    for ax, (m, lab) in zip(axes.flat, panels):
        for rule in ("baseline", "timed"):
            g = df[df.rule == rule]
            ax.plot(g[xcol], g[m], marker="o", color=COLORS[rule], label=LABELS[rule])
        if vline:
            ax.axvline(vline, ls="--", color="gray", lw=1, label=vlabel)
        ax.set_xlabel(xlabel)
        ax.set_ylabel(lab)
        ax.set_title(lab)
    axes[0, 0].legend(fontsize=8)
    fig.suptitle(title)
    fig.tight_layout()
    fig.savefig(os.path.join(OUT, fname), dpi=150)
    plt.close(fig)


def assumptions_table():
    c = config
    rows = [
        ("Terminal / route", "Laoag City terminal; Laoag-Batac-Paoay round trip", "Reported"),
        ("Fleet size N", c.FLEET_SIZES[0], "Reported (15 jeepneys)"),
        ("Seat capacity C", c.CAPACITY, "Reported (about 20 when full)"),
        ("Timed rule T (min)", c.TIME_LIMITS[0], "Reported (leave after about 5 min)"),
        ("Peak arrivals (passengers/min)", 1 / c.ARRIVAL_MEAN_MIN["peak"], "ASSUMPTION (survey unit unclear)"),
        ("Off-peak arrivals (passengers/min)", 1 / c.ARRIVAL_MEAN_MIN["offpeak"], "ASSUMPTION (survey unit unclear)"),
        ("Round trip, peak (min, mean)", c.TRIP_MEAN_MIN["peak"], "Reported 60 min heavy traffic; peak = heavy is ASSUMPTION"),
        ("Round trip, off-peak (min, mean)", c.TRIP_MEAN_MIN["offpeak"], "Reported 40 min normal traffic"),
        ("Round trip SD (min)", c.TRIP_SD_MIN, "ASSUMPTION"),
        ("Boarding time mean / SD (sec)", f"{c.BOARD_MEAN_SEC} / {c.BOARD_SD_SEC}", "ASSUMPTION"),
        ("Boarding minimum (sec)", c.BOARD_MIN_SEC, "ASSUMPTION"),
        ("Warm-up (min)", c.WARMUP_MIN, "ASSUMPTION"),
        ("Measured run length (min)", c.RUN_LENGTH_MIN, "ASSUMPTION"),
        ("Replications per scenario", c.REPLICATIONS, "Design choice"),
    ]
    return pd.DataFrame(rows, columns=["Parameter", "Value", "Source"])


def pm(mean, hw, digits=2):
    return f"{mean:.{digits}f} +/- {hw:.{digits}f}"


def summary_for_slides(summ):
    out = pd.DataFrame({"Scenario": summ.scenario, "Period": summ.period, "Rule": summ.rule})
    for m, lab in MEASURES.items():
        out[lab] = [pm(a, b) for a, b in zip(summ[m + "_mean"], summ[m + "_hw"])]
    return out


def format_sheets(path):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F5C1F")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:60])
            ws.column_dimensions[col[0].column_letter].width = min(max(12, width + 2), 45)
        ws.freeze_panes = "B2"
    wb.save(path)


if __name__ == "__main__":
    df = pd.read_csv(os.path.join(OUT, "results.csv"))
    summ = summarize(df)
    diff = pd.read_csv(os.path.join(OUT, "diff_vs_baseline.csv"))
    arr = pd.read_csv(os.path.join(OUT, "sensitivity_arrival.csv"))
    fleet = pd.read_csv(os.path.join(OUT, "sensitivity_fleet.csv"))

    plot_key_measures(summ)
    cap = capacity_per_min(config.FLEET_SIZES[0])
    plot_sweep(arr, "arrival_per_min", "Passenger arrival rate (passengers per minute)",
               f"Effect of arrival rate, N = {config.FLEET_SIZES[0]}", "sensitivity_arrival.png",
               cap, f"approx. fleet capacity ({cap:.1f}/min)")
    plot_sweep(fleet, "fleet_size", "Fleet size N (jeepneys)",
               "Effect of fleet size at the peak arrival rate", "sensitivity_fleet.png")

    path = os.path.join(OUT, "results_report.xlsx")
    with pd.ExcelWriter(path) as xw:
        assumptions_table().to_excel(xw, sheet_name="Assumptions", index=False)
        summary_for_slides(summ).to_excel(xw, sheet_name="Summary (mean, 95% CI)", index=False)
        summ.to_excel(xw, sheet_name="Summary (numeric)", index=False)
        diff.to_excel(xw, sheet_name="Timed minus baseline", index=False)
        arr.round(3).to_excel(xw, sheet_name="Sweep arrival rate", index=False)
        fleet.round(3).to_excel(xw, sheet_name="Sweep fleet size", index=False)
        df.round(4).to_excel(xw, sheet_name="Raw replications", index=False)
    format_sheets(path)
    print(f"saved {path} and charts in {OUT}")
